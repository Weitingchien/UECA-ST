#!/usr/bin/env python3
"""Export selected 10-fold cumulative micro metrics to an Excel workbook."""

from __future__ import annotations

import argparse
import re
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


REPO_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_RESULTS_DIR = Path("results_ep_split10_t1v1te1_u7_aligned_disjoint_2019")
DEFAULT_OUTPUT_NAME = "method_micro_metrics_summary.xlsx"
DEFAULT_EXPERIMENT_ROOTS = (
    Path("H:/ep_split10_t1v1te1_u7_aligned_disjoint_2019"),
    Path("I:/ep_split10_t1v1te1_u7_aligned_disjoint_2019"),
    Path("F:/ep_split10_t1v1te1_u7_aligned_disjoint_2019"),
)
COUNTS_FILENAMES = ("counts_summary_detailed.txt", "counts_summary_detailed")

TASKS = ("Emotion", "Cause", "Pair (m1)", "Pair (m2)", "Pair (m3)")
METRICS = ("P", "R", "F1")

SUMMARY_HEADER_RE = re.compile(
    r"^\s*10\s*折累計統計結果\s*\(Mean\s*±\s*Std\)\s*:\s*$"
)
SECTION_RE = re.compile(r"^\s*(Emotion|Cause|Pair\s*\(m[123]\))\s*:\s*$")
METRIC_RE = re.compile(
    r"^\s*(P|R|F1)\s*\(micro\)\s*:\s*"
    r"(\d+(?:\.\d+)?)%\s*±\s*(\d+(?:\.\d+)?)%\s*$"
)
BLOCK_END_RE = re.compile(r"^\s*=+\s*$")
EXPERIMENTS_HEADER_RE = re.compile(
    r"^\s*彙整的實驗目錄\s*\(共\s*(\d+)\s*個\)\s*[:：]\s*$"
)
EXPERIMENT_LINE_RE = re.compile(r"^\s*-\s+(.+?)\s*$")
SEED_RE = re.compile(r"_seed(\d+)(?:_|$)")
COUNTS_HEADER_RE = re.compile(r"^\s*10\s*折累計統計結果\s*[:：]\s*$")
COUNT_RE = re.compile(
    r"^\s*Total\s+(TP|Pred|GT)\b[^:：]*[:：]\s*(\d+)\s*$"
)
COUNTS_RATE_RE = re.compile(
    r"^\s*P\s*\(micro\)\s*[:：]\s*(\d+(?:\.\d+)?)%\s*\|\s*"
    r"R\s*\(micro\)\s*[:：]\s*(\d+(?:\.\d+)?)%\s*\|\s*"
    r"F1\s*\(micro\)\s*[:：]\s*(\d+(?:\.\d+)?)%\s*$"
)

TITLE_FILL = PatternFill("solid", fgColor="17365D")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
SUBTITLE_FILL = PatternFill("solid", fgColor="DDEBF7")
WHITE_FONT = Font(color="FFFFFF", bold=True)
THIN_GRAY_BOTTOM = Border(bottom=Side(style="thin", color="A6A6A6"))


@dataclass(frozen=True)
class MetricValue:
    mean: float
    std: float


@dataclass(frozen=True)
class CountsValue:
    total_tp: int
    total_pred: int
    total_gt: int
    p_micro: float
    r_micro: float
    f1_micro: float


@dataclass(frozen=True)
class ExperimentCounts:
    order: int
    method: str
    seed: int
    experiment_name: str
    experiment_root: Path
    counts_file: Path
    values: dict[str, CountsValue]


@dataclass(frozen=True)
class MethodSpec:
    order: int
    method: str
    summary_file: str
    method_family: str
    selection_basis: str
    training_target: str
    final_model: str
    beta: str
    note: str = ""


METHOD_SPECS = (
    MethodSpec(
        1,
        "-Mconf-Semo-Tpair-Fbest",
        "UECA_Prompt_f1-10_i70_lr1e-5_bs8_wd0.01_confidence_emotion_clause_"
        "confmul1_gamma0.5_nlmnest_st5_ste20_remove_pseudo_v2_CE_summary.txt",
        "信心程度篩選（Mconf）",
        "情緒子句（Semo）",
        "Pair（Tpair）",
        "單一任務最佳模型（Fbest）",
        "不適用",
        "依方法命名規則修正原清單中與第 6 項重複的檔名。",
    ),
    MethodSpec(
        2,
        "-Mconf-Semo-Tpair-Favg",
        "UECA_Prompt_f1-10_i70_lr1e-5_bs8_wd0.01_confidence_emotion_clause_"
        "confmul1_gamma0.5_nlmnest_avgs_simple_st5_ste20_remove_pseudo_v2_CE_summary.txt",
        "信心程度篩選（Mconf）",
        "情緒子句（Semo）",
        "Pair（Tpair）",
        "參數平均（Favg）",
        "不適用",
    ),
    MethodSpec(
        3,
        "-Mours-Semo-Tpair-Fbest-Beta0.1",
        "UECA_Prompt_f1-10_i70_lr1e-5_bs8_wd0.01_nest_k5_emotion_clause_"
        "knnemotion_clause_nestmul1_nbeta0.1_nm0.6_gamma0.5_nlmnest_st5_ste20_"
        "remove_pseudo_v2_CE_summary.txt",
        "本研究方法（Mours）",
        "情緒子句（Semo）",
        "Pair（Tpair）",
        "單一任務最佳模型（Fbest）",
        "0.1",
    ),
    MethodSpec(
        4,
        "-Mours-Semo-Tpair-Favg-Beta0.1",
        "UECA_Prompt_f1-10_i70_lr1e-5_bs8_wd0.01_nest_k5_emotion_clause_"
        "knnemotion_clause_nestmul1_nbeta0.1_nm0.6_gamma0.5_nlmnest_avgs_simple_"
        "st5_ste20_remove_pseudo_v2_CE_summary.txt",
        "本研究方法（Mours）",
        "情緒子句（Semo）",
        "Pair（Tpair）",
        "參數平均（Favg）",
        "0.1",
    ),
    MethodSpec(
        5,
        "-Mconf-Scau-Tpair-Fbest",
        "UECA_Prompt_f1-10_i70_lr1e-5_bs8_wd0.01_th0.9_maskemotion_gamma0.5_"
        "st5_ste20_remove_pseudo_v2_CE_summary.txt",
        "信心程度篩選（Mconf）",
        "原因子句（Scau）",
        "Pair（Tpair）",
        "單一任務最佳模型（Fbest）",
        "不適用",
        "依使用者指定對應 th0.9_maskemotion 實驗。",
    ),
    MethodSpec(
        6,
        "-Mconf-Scau-Tpair-Favg",
        "UECA_Prompt_f1-10_i70_lr1e-5_bs8_wd0.01_confidence_cause_clause_"
        "confmul1_gamma0.5_nlmnest_avgs_simple_st5_ste20_remove_pseudo_v2_CE_summary.txt",
        "信心程度篩選（Mconf）",
        "原因子句（Scau）",
        "Pair（Tpair）",
        "參數平均（Favg）",
        "不適用",
    ),
    MethodSpec(
        7,
        "-Mours-Scau-Tpair-Fbest-Beta0.1",
        "UECA_Prompt_f1-10_i70_lr1e-5_bs8_wd0.01_nest_k5_cause_clause_"
        "knncause_clause_nestmul1_nbeta0.1_nm0.6_gamma0.5_nlmnest_st5_ste20_"
        "remove_pseudo_v2_CE_summary.txt",
        "本研究方法（Mours）",
        "原因子句（Scau）",
        "Pair（Tpair）",
        "單一任務最佳模型（Fbest）",
        "0.1",
    ),
    MethodSpec(
        8,
        "-Mours-Scau-Tpair-Favg-Beta0.1",
        "UECA_Prompt_f1-10_i70_lr1e-5_bs8_wd0.01_nest_k5_cause_clause_"
        "knncause_clause_nestmul1_nbeta0.1_nm0.6_gamma0.5_nlmnest_avgs_simple_"
        "st5_ste20_remove_pseudo_v2_CE_summary.txt",
        "本研究方法（Mours）",
        "原因子句（Scau）",
        "Pair（Tpair）",
        "參數平均（Favg）",
        "0.1",
    ),
    MethodSpec(
        9,
        "-Mours-Semo-Temo-Fbest-Beta0.1",
        "UECA_Prompt_f1-10_i70_lr1e-5_bs8_wd0.01_nest_k5_emotion_clause_"
        "knnemotion_clause_nestmul1_nbeta0.1_nm0.6_gamma0.5_nlmnest_st5_ste20_"
        "initemo_stemo_testemo_remove_pseudo_v2_CE_summary.txt",
        "本研究方法（Mours）",
        "情緒子句（Semo）",
        "Emotion（Temo）",
        "單一任務最佳模型（Fbest）",
        "0.1",
    ),
    MethodSpec(
        10,
        "-Mours-Scau-Tcau-Fbest-Beta0.1",
        "UECA_Prompt_f1-10_i70_lr1e-5_bs8_wd0.01_nest_k5_cause_clause_"
        "knncause_clause_nestmul1_nbeta0.1_nm0.6_gamma0.5_nlmnest_st5_ste20_"
        "initcau_stcau_testcau_remove_pseudo_v2_CE_summary.txt",
        "本研究方法（Mours）",
        "原因子句（Scau）",
        "Cause（Tcau）",
        "單一任務最佳模型（Fbest）",
        "0.1",
    ),
)


NAMING_RULES = (
    ("-M", "名稱前綴", "方法名稱", "表示後續字串為實驗方法名稱。"),
    ("-Mconf", "方法類別", "confidence_*；第 5 項為 th0.9_maskemotion", "根據信心程度篩選。"),
    ("-Mours", "方法類別", "nest_k5_*", "本研究提出的方法。"),
    ("-S", "自訓練元件", "nest / knn / divergence 相關設定", "表示自訓練時會計算 KNN 與散度。"),
    ("-Semo", "篩選／表徵依據", "emotion_clause / knnemotion_clause", "以情緒子句嵌入計算 KNN，並以情緒子句標籤分布計算散度。"),
    ("-Scau", "篩選／表徵依據", "cause_clause / knncause_clause", "以原因子句嵌入與原因子句標籤分布作為對應依據。"),
    ("-Tpair", "最佳模型任務", "一般 st5_ste20 流程", "每輪自訓練選用 Pair 表現最佳的模型。"),
    ("-Temo", "最佳模型任務", "initemo_stemo_testemo", "每輪自訓練選用 Emotion 表現最佳的模型。"),
    ("-Tcau", "最佳模型任務", "initcau_stcau_testcau", "每輪自訓練選用 Cause 表現最佳的模型。"),
    ("-Favg", "最終模型", "nlmnest_avgs_simple", "自訓練結束後進行參數平均。"),
    ("-Fbest", "最終模型", "nlmnest_st5 或未含 avgs_simple", "使用對應單一任務的最佳模型作為最終模型。"),
    ("-Beta{x}", "權重", "nbeta{x}", "KL_L(Nj) 的權重為 x。"),
    ("-Beta0", "權重", "nbeta0", "不考慮 KNN 鄰居彼此的標籤分布差異。"),
)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="整理指定方法的 10 折累計 micro P/R/F1，輸出為 Excel。"
    )
    parser.add_argument(
        "--results-dir",
        default=str(DEFAULT_RESULTS_DIR),
        help="包含 *_summary.txt 的資料夾（相對路徑以專案根目錄為基準）。",
    )
    parser.add_argument(
        "--output",
        default=None,
        help="輸出 .xlsx 路徑；省略時寫入 results-dir/method_micro_metrics_summary.xlsx。",
    )
    parser.add_argument(
        "--experiment-roots",
        nargs="+",
        default=[str(path) for path in DEFAULT_EXPERIMENT_ROOTS],
        help="依序搜尋實驗目錄的根路徑；預設順序為 H、I、F。",
    )
    return parser.parse_args()


def resolve_repo_path(value: str | Path) -> Path:
    path = Path(value)
    return path if path.is_absolute() else REPO_ROOT / path


def parse_summary(path: Path) -> dict[tuple[str, str], MetricValue]:
    lines = path.read_text(encoding="utf-8").splitlines()
    header_indexes = [
        index for index, line in enumerate(lines) if SUMMARY_HEADER_RE.fullmatch(line)
    ]
    if len(header_indexes) != 1:
        raise ValueError(
            f"{path.name}: 預期恰有一個『10 折累計統計結果 (Mean ± Std)』區塊，"
            f"實際找到 {len(header_indexes)} 個。"
        )

    values: dict[tuple[str, str], MetricValue] = {}
    current_task: str | None = None
    for line in lines[header_indexes[0] + 1 :]:
        if BLOCK_END_RE.fullmatch(line):
            break

        section_match = SECTION_RE.fullmatch(line)
        if section_match:
            current_task = re.sub(r"\s+", " ", section_match.group(1))
            continue

        metric_match = METRIC_RE.fullmatch(line)
        if not metric_match:
            continue
        if current_task is None:
            raise ValueError(f"{path.name}: 指標列之前缺少任務標題：{line.strip()}")

        metric, mean_text, std_text = metric_match.groups()
        key = (current_task, metric)
        if key in values:
            raise ValueError(f"{path.name}: 重複指標 {current_task} / {metric}。")

        mean = float(mean_text) / 100.0
        std = float(std_text) / 100.0
        if not (0.0 <= mean <= 1.0 and 0.0 <= std <= 1.0):
            raise ValueError(f"{path.name}: {current_task} / {metric} 超出 0%–100%。")
        values[key] = MetricValue(mean=mean, std=std)

    expected = {(task, metric) for task in TASKS for metric in METRICS}
    if values.keys() != expected:
        missing = sorted(expected - values.keys())
        extra = sorted(values.keys() - expected)
        raise ValueError(f"{path.name}: 指標欄位不完整；缺少={missing}，多出={extra}。")
    return values


def parse_experiment_names(summary_path: Path) -> tuple[str, ...]:
    lines = summary_path.read_text(encoding="utf-8").splitlines()
    headers = [
        (index, match)
        for index, line in enumerate(lines)
        if (match := EXPERIMENTS_HEADER_RE.fullmatch(line))
    ]
    if len(headers) != 1:
        raise ValueError(
            f"{summary_path.name}: 預期恰有一個『彙整的實驗目錄』區塊，"
            f"實際找到 {len(headers)} 個。"
        )

    header_index, header_match = headers[0]
    expected_count = int(header_match.group(1))
    names: list[str] = []
    for line in lines[header_index + 1 :]:
        if not line.strip():
            if names:
                break
            continue
        match = EXPERIMENT_LINE_RE.fullmatch(line)
        if match:
            names.append(match.group(1))
        elif names:
            break

    if len(names) != expected_count:
        raise ValueError(
            f"{summary_path.name}: 標示 {expected_count} 個實驗目錄，實際解析到 {len(names)} 個。"
        )
    if len(names) != len(set(names)):
        raise ValueError(f"{summary_path.name}: 實驗目錄名稱重複。")
    return tuple(names)


def extract_seed(experiment_name: str) -> int:
    match = SEED_RE.search(experiment_name)
    if match is None:
        raise ValueError(f"實驗目錄名稱缺少 seed：{experiment_name}")
    return int(match.group(1))


def locate_counts_file(
    experiment_name: str, experiment_roots: tuple[Path, ...]
) -> tuple[Path, Path]:
    matches: list[tuple[Path, Path]] = []
    attempted: list[Path] = []
    for root in experiment_roots:
        experiment_dir = root / experiment_name
        for filename in COUNTS_FILENAMES:
            candidate = experiment_dir / filename
            attempted.append(candidate)
            if candidate.is_file():
                matches.append((root, candidate))
                break

    if not matches:
        tried = "\n  - ".join(str(path) for path in attempted)
        raise FileNotFoundError(
            f"找不到 {experiment_name} 的 counts_summary_detailed；已嘗試：\n  - {tried}"
        )
    if len(matches) > 1:
        found = "\n  - ".join(str(path) for _, path in matches)
        raise ValueError(f"跨根目錄找到重複的 counts 檔案：\n  - {found}")
    return matches[0]


def parse_counts_file(path: Path) -> dict[str, CountsValue]:
    lines = path.read_text(encoding="utf-8").splitlines()
    header_indexes = [
        index for index, line in enumerate(lines) if COUNTS_HEADER_RE.fullmatch(line)
    ]
    if not header_indexes:
        raise ValueError(f"{path}: 找不到『10 折累計統計結果』區塊。")

    totals: dict[tuple[str, str], int] = {}
    rates: dict[tuple[str, str], float] = {}
    current_task: str | None = None
    for line in lines[header_indexes[-1] + 1 :]:
        if BLOCK_END_RE.fullmatch(line):
            break

        section_match = SECTION_RE.fullmatch(line)
        if section_match:
            current_task = re.sub(r"\s+", " ", section_match.group(1))
            continue

        count_match = COUNT_RE.fullmatch(line)
        if count_match:
            if current_task is None:
                raise ValueError(f"{path}: Total 欄位之前缺少任務標題。")
            count_name, count_text = count_match.groups()
            key = (current_task, count_name)
            if key in totals:
                raise ValueError(f"{path}: 重複計數 {current_task} / Total {count_name}。")
            totals[key] = int(count_text)
            continue

        rate_match = COUNTS_RATE_RE.fullmatch(line)
        if rate_match:
            if current_task is None:
                raise ValueError(f"{path}: micro 指標之前缺少任務標題。")
            for metric, value in zip(METRICS, rate_match.groups()):
                key = (current_task, metric)
                if key in rates:
                    raise ValueError(f"{path}: 重複指標 {current_task} / {metric}。")
                rates[key] = float(value) / 100.0

    expected_totals = {
        (task, count_name)
        for task in TASKS
        for count_name in ("TP", "Pred", "GT")
    }
    expected_rates = {(task, metric) for task in TASKS for metric in METRICS}
    if totals.keys() != expected_totals:
        missing = sorted(expected_totals - totals.keys())
        extra = sorted(totals.keys() - expected_totals)
        raise ValueError(f"{path}: Total 欄位不完整；缺少={missing}，多出={extra}。")
    if rates.keys() != expected_rates:
        missing = sorted(expected_rates - rates.keys())
        extra = sorted(rates.keys() - expected_rates)
        raise ValueError(f"{path}: micro 指標不完整；缺少={missing}，多出={extra}。")

    values: dict[str, CountsValue] = {}
    for task in TASKS:
        total_tp = totals[(task, "TP")]
        total_pred = totals[(task, "Pred")]
        total_gt = totals[(task, "GT")]
        if total_pred <= 0 or total_gt <= 0:
            raise ValueError(f"{path}: {task} 的 Total Pred/GT 必須大於 0。")
        if total_tp < 0 or total_tp > min(total_pred, total_gt):
            raise ValueError(f"{path}: {task} 的 Total TP 不合理。")

        calculated = {
            "P": total_tp / total_pred,
            "R": total_tp / total_gt,
            "F1": 2 * total_tp / (total_pred + total_gt),
        }
        for metric in METRICS:
            if abs(calculated[metric] - rates[(task, metric)]) > 0.000051:
                raise ValueError(f"{path}: {task} / {metric} 與 Total 計數不一致。")

        values[task] = CountsValue(
            total_tp=total_tp,
            total_pred=total_pred,
            total_gt=total_gt,
            p_micro=rates[(task, "P")],
            r_micro=rates[(task, "R")],
            f1_micro=rates[(task, "F1")],
        )
    return values


def collect_experiment_counts(
    results_dir: Path, experiment_roots: tuple[Path, ...]
) -> list[ExperimentCounts]:
    records: list[ExperimentCounts] = []
    for spec in METHOD_SPECS:
        summary_path = results_dir / spec.summary_file
        experiment_names = parse_experiment_names(summary_path)
        method_records: list[ExperimentCounts] = []
        for experiment_name in experiment_names:
            experiment_root, counts_file = locate_counts_file(
                experiment_name, experiment_roots
            )
            method_records.append(
                ExperimentCounts(
                    order=spec.order,
                    method=spec.method,
                    seed=extract_seed(experiment_name),
                    experiment_name=experiment_name,
                    experiment_root=experiment_root,
                    counts_file=counts_file,
                    values=parse_counts_file(counts_file),
                )
            )

        seeds = [record.seed for record in method_records]
        if len(seeds) != len(set(seeds)):
            raise ValueError(f"{summary_path.name}: seed 重複：{seeds}")
        records.extend(sorted(method_records, key=lambda record: record.seed))
    return records


def style_title(sheet, title: str, subtitle: str, last_column: int) -> None:
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    title_cell = sheet.cell(1, 1, title)
    title_cell.fill = TITLE_FILL
    title_cell.font = Font(color="FFFFFF", bold=True, size=16)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.row_dimensions[1].height = 28

    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_column)
    subtitle_cell = sheet.cell(2, 1, subtitle)
    subtitle_cell.fill = SUBTITLE_FILL
    subtitle_cell.font = Font(color="1F1F1F", italic=True, size=10)
    subtitle_cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[2].height = 30


def style_header(sheet, row: int, last_column: int) -> None:
    for cell in sheet[row][:last_column]:
        cell.fill = HEADER_FILL
        cell.font = WHITE_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_GRAY_BOTTOM
    sheet.row_dimensions[row].height = 32


def add_table(sheet, name: str, reference: str) -> None:
    table = Table(displayName=name, ref=reference)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)


def configure_print(sheet, print_area: str, repeat_rows: str = "1:4") -> None:
    sheet.sheet_view.showGridLines = False
    sheet.print_area = print_area
    sheet.print_title_rows = repeat_rows
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.sheet_properties.pageSetUpPr.autoPageBreaks = False
    sheet.page_margins.left = 0.25
    sheet.page_margins.right = 0.25
    sheet.page_margins.top = 0.5
    sheet.page_margins.bottom = 0.5


def create_metrics_sheet(workbook: Workbook, parsed_by_method) -> None:
    sheet = workbook.active
    sheet.title = "Micro Metrics"
    style_title(
        sheet,
        "10 折累計 Micro 指標彙整",
        "來源：各 summary 的 3 個 seed 百分比算術平均與樣本標準差；原始分子／分母請見 Seed Counts，三個 seed 合計請見 Pooled Counts。",
        9,
    )
    headers = (
        "序號",
        "方法名稱",
        "任務",
        "P (micro) Mean",
        "P (micro) Std",
        "R (micro) Mean",
        "R (micro) Std",
        "F1 (micro) Mean",
        "F1 (micro) Std",
    )
    sheet.append([])
    sheet.append(headers)

    for spec in METHOD_SPECS:
        values = parsed_by_method[spec.method]
        for task in TASKS:
            sheet.append(
                (
                    spec.order,
                    spec.method,
                    task,
                    values[(task, "P")].mean,
                    values[(task, "P")].std,
                    values[(task, "R")].mean,
                    values[(task, "R")].std,
                    values[(task, "F1")].mean,
                    values[(task, "F1")].std,
                )
            )

    style_header(sheet, 4, 9)
    for row in sheet.iter_rows(min_row=5, max_row=sheet.max_row):
        row[0].alignment = Alignment(horizontal="center", vertical="center")
        row[1].alignment = Alignment(vertical="center", wrap_text=True)
        row[2].alignment = Alignment(horizontal="center", vertical="center")
        for cell in row[3:]:
            cell.number_format = "0.00%"
            cell.alignment = Alignment(horizontal="center", vertical="center")

    widths = {"A": 8, "B": 43, "C": 14, "D": 17, "E": 16, "F": 17, "G": 16, "H": 18, "I": 17}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "C5"
    add_table(sheet, "MicroMetricsTable", f"A4:I{sheet.max_row}")
    configure_print(sheet, f"A1:I{sheet.max_row}")


def create_seed_counts_sheet(
    workbook: Workbook, records: list[ExperimentCounts]
) -> None:
    sheet = workbook.create_sheet("Seed Counts")
    style_title(
        sheet,
        "各 Seed 的 10 折累計計數",
        "每列取自單一 counts_summary_detailed；P = TP / Pred，R = TP / GT，F1 = 2TP / (Pred + GT)。來源根目錄依 H → I → F 搜尋。",
        13,
    )
    headers = (
        "序號",
        "方法名稱",
        "Seed",
        "任務",
        "Total TP",
        "Total Pred",
        "Total GT",
        "P (micro)",
        "R (micro)",
        "F1 (micro)",
        "來源根目錄",
        "實驗目錄",
        "Counts 檔案",
    )
    sheet.append([])
    sheet.append(headers)
    for record in records:
        for task in TASKS:
            counts = record.values[task]
            sheet.append(
                (
                    record.order,
                    record.method,
                    record.seed,
                    task,
                    counts.total_tp,
                    counts.total_pred,
                    counts.total_gt,
                    counts.p_micro,
                    counts.r_micro,
                    counts.f1_micro,
                    str(record.experiment_root),
                    record.experiment_name,
                    record.counts_file.name,
                )
            )

    style_header(sheet, 4, 13)
    for row in sheet.iter_rows(min_row=5, max_row=sheet.max_row):
        row[0].alignment = Alignment(horizontal="center", vertical="center")
        row[1].alignment = Alignment(vertical="center", wrap_text=True)
        for cell in (row[2], row[3]):
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for cell in row[4:7]:
            cell.number_format = "#,##0"
            cell.alignment = Alignment(horizontal="right", vertical="center")
        for cell in row[7:10]:
            cell.number_format = "0.00%"
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for cell in row[10:]:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.row_dimensions[row[0].row].height = 52

    widths = {
        "A": 8,
        "B": 43,
        "C": 10,
        "D": 14,
        "E": 14,
        "F": 15,
        "G": 14,
        "H": 14,
        "I": 14,
        "J": 15,
        "K": 47,
        "L": 82,
        "M": 30,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "E5"
    add_table(sheet, "SeedCountsTable", f"A4:M{sheet.max_row}")
    configure_print(sheet, f"A1:M{sheet.max_row}")


def create_pooled_counts_sheet(
    workbook: Workbook, record_count: int
) -> None:
    sheet = workbook.create_sheet("Pooled Counts", 1)
    style_title(
        sheet,
        "三個 Seed 的 Pooled Counts",
        "Total 為三個 seed（共 30 次 fold 評估）之合計，同一測試資料會重複計入 3 次；Pooled 比例不等同於 Micro Metrics 的 seed 百分比算術平均。",
        10,
    )
    headers = (
        "序號",
        "方法名稱",
        "任務",
        "Seed 數",
        "Total TP",
        "Total Pred",
        "Total GT",
        "Pooled P",
        "Pooled R",
        "Pooled F1",
    )
    sheet.append([])
    sheet.append(headers)

    seed_last_row = 4 + record_count * len(TASKS)
    method_range = f"'Seed Counts'!$B$5:$B${seed_last_row}"
    task_range = f"'Seed Counts'!$D$5:$D${seed_last_row}"
    for spec in METHOD_SPECS:
        for task in TASKS:
            row_number = sheet.max_row + 1
            method_criterion = f"$B{row_number}"
            task_criterion = f"$C{row_number}"
            sheet.append(
                (
                    spec.order,
                    spec.method,
                    task,
                    f"=COUNTIFS({method_range},{method_criterion},{task_range},{task_criterion})",
                    f"=SUMIFS('Seed Counts'!$E$5:$E${seed_last_row},{method_range},{method_criterion},{task_range},{task_criterion})",
                    f"=SUMIFS('Seed Counts'!$F$5:$F${seed_last_row},{method_range},{method_criterion},{task_range},{task_criterion})",
                    f"=SUMIFS('Seed Counts'!$G$5:$G${seed_last_row},{method_range},{method_criterion},{task_range},{task_criterion})",
                    f"=IFERROR(E{row_number}/F{row_number},0)",
                    f"=IFERROR(E{row_number}/G{row_number},0)",
                    f"=IFERROR(2*E{row_number}/(F{row_number}+G{row_number}),0)",
                )
            )

    style_header(sheet, 4, 10)
    for row in sheet.iter_rows(min_row=5, max_row=sheet.max_row):
        row[0].alignment = Alignment(horizontal="center", vertical="center")
        row[1].alignment = Alignment(vertical="center", wrap_text=True)
        row[2].alignment = Alignment(horizontal="center", vertical="center")
        for cell in row[3:7]:
            cell.number_format = "#,##0"
            cell.alignment = Alignment(horizontal="right", vertical="center")
        for cell in row[7:]:
            cell.number_format = "0.00%"
            cell.alignment = Alignment(horizontal="center", vertical="center")

    widths = {
        "A": 8,
        "B": 43,
        "C": 14,
        "D": 12,
        "E": 15,
        "F": 16,
        "G": 15,
        "H": 15,
        "I": 15,
        "J": 16,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "C5"
    add_table(sheet, "PooledCountsTable", f"A4:J{sheet.max_row}")
    configure_print(sheet, f"A1:J{sheet.max_row}")


def create_mapping_sheet(workbook: Workbook) -> None:
    sheet = workbook.create_sheet("Method Mapping")
    style_title(
        sheet,
        "方法名稱與 Summary 檔案對照",
        "第 1 項依 Semo／Fbest 命名規則修正為 confidence_emotion_clause + nlmnest_st5，避免與第 6 項重複。",
        9,
    )
    headers = (
        "序號",
        "方法名稱",
        "Summary 檔名",
        "方法類別",
        "篩選／表徵依據",
        "最佳模型任務",
        "最終模型",
        "Beta",
        "備註",
    )
    sheet.append([])
    sheet.append(headers)
    for spec in METHOD_SPECS:
        sheet.append(
            (
                spec.order,
                spec.method,
                spec.summary_file,
                spec.method_family,
                spec.selection_basis,
                spec.training_target,
                spec.final_model,
                spec.beta,
                spec.note,
            )
        )

    style_header(sheet, 4, 9)
    for row in sheet.iter_rows(min_row=5, max_row=sheet.max_row):
        row[0].alignment = Alignment(horizontal="center", vertical="center")
        for cell in row[1:]:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        row[7].alignment = Alignment(horizontal="center", vertical="center")
        sheet.row_dimensions[row[0].row].height = 64

    widths = {"A": 8, "B": 39, "C": 68, "D": 23, "E": 22, "F": 21, "G": 27, "H": 12, "I": 38}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "C5"
    add_table(sheet, "MethodMappingTable", f"A4:I{sheet.max_row}")
    configure_print(sheet, f"A1:I{sheet.max_row}")


def create_naming_rules_sheet(workbook: Workbook) -> None:
    sheet = workbook.create_sheet("Naming Rules")
    style_title(
        sheet,
        "實驗方法命名規則",
        "論文中的方法標記與 summary 檔名線索對照；實際方法名稱可依序串接各類別標記。",
        4,
    )
    headers = ("論文名稱標記", "類別", "Summary 檔名線索", "定義")
    sheet.append([])
    sheet.append(headers)
    for rule in NAMING_RULES:
        sheet.append(rule)

    style_header(sheet, 4, 4)
    for row in sheet.iter_rows(min_row=5, max_row=sheet.max_row):
        row[0].font = Font(bold=True, color="17365D")
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.row_dimensions[row[0].row].height = 42

    widths = {"A": 20, "B": 20, "C": 43, "D": 75}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A5"
    add_table(sheet, "NamingRulesTable", f"A4:D{sheet.max_row}")
    configure_print(sheet, f"A1:D{sheet.max_row}")


def build_workbook(
    results_dir: Path,
    experiment_roots: tuple[Path, ...] = DEFAULT_EXPERIMENT_ROOTS,
) -> Workbook:
    parsed_by_method: dict[str, dict[tuple[str, str], MetricValue]] = {}
    for spec in METHOD_SPECS:
        summary_path = results_dir / spec.summary_file
        if not summary_path.is_file():
            raise FileNotFoundError(f"找不到 summary 檔案：{summary_path}")
        parsed_by_method[spec.method] = parse_summary(summary_path)
    counts_records = collect_experiment_counts(results_dir, experiment_roots)

    workbook = Workbook()
    workbook.properties.creator = "UECA_ST"
    workbook.properties.title = "10 折累計 Micro 指標與實際計數彙整"
    workbook.properties.subject = "10 個實驗方法的 micro P/R/F1 與 Total TP/Pred/GT"
    workbook.properties.description = (
        "百分比取自 summary 的 Mean ± Std；計數取自各 seed 的 counts_summary_detailed。"
    )
    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    create_metrics_sheet(workbook, parsed_by_method)
    create_seed_counts_sheet(workbook, counts_records)
    create_pooled_counts_sheet(workbook, len(counts_records))
    create_mapping_sheet(workbook)
    create_naming_rules_sheet(workbook)
    return workbook


def main() -> None:
    args = parse_args()
    results_dir = resolve_repo_path(args.results_dir)
    if not results_dir.is_dir():
        raise NotADirectoryError(f"results-dir 不存在：{results_dir}")

    output = (
        resolve_repo_path(args.output)
        if args.output
        else results_dir / DEFAULT_OUTPUT_NAME
    )
    if output.suffix.lower() != ".xlsx":
        raise ValueError(f"輸出檔案必須是 .xlsx：{output}")

    experiment_roots = tuple(resolve_repo_path(path) for path in args.experiment_roots)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = build_workbook(results_dir, experiment_roots)
    workbook.save(output)
    print(f"已輸出：{output}")


if __name__ == "__main__":
    main()
